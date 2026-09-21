"""Eenmalige import van de catalogus-tabbladen uit het stamgegevens-sheet
(Google Sheets, geexporteerd als xlsx). Bestaande liederen blijven staan;
alleen nieuwe id's worden toegevoegd."""

from __future__ import annotations

import re
from pathlib import Path

import openpyxl

from .catalogus import Catalogus
from .modellen import Lied, Referentie
from .tekst import slug

TPP_TITEL = re.compile(r"^ps\.?\s*(\d+)\s*[-–]\s*(.+)$", re.IGNORECASE)


def importeer_sheet(xlsx: Path, catalogus: Catalogus) -> dict[str, int]:
    wb = openpyxl.load_workbook(xlsx, read_only=True, data_only=True)
    toegevoegd = {"opw.yaml": 0, "sela.yaml": 0, "tpp.yaml": 0, "svg.yaml": 0}

    for nr, titel in _rijen(wb, "Bron Opwekking", (1, 2)):
        if not str(nr).strip().isdigit():
            continue
        delen = [d.strip() for d in str(titel).split("/") if d.strip()]
        lied = Lied(id=f"opw-{int(nr)}", titel=delen[0] if delen else f"Opwekking {nr}", referenties=[Referentie("opw", str(int(nr)))], aliassen=delen[1:])
        toegevoegd["opw.yaml"] += _voeg_toe(catalogus, lied, "opw.yaml")

    for (titel,) in _rijen(wb, "Bron Sela", (1,)):
        lied = Lied(id=f"sela-{slug(titel)}", titel=str(titel).strip(), artiest="Sela")
        toegevoegd["sela.yaml"] += _voeg_toe(catalogus, lied, "sela.yaml")

    for (tekst,) in _rijen(wb, "Bron TPP", (1,)):
        m = TPP_TITEL.match(str(tekst).strip())
        if m:
            nummer, titel = m.group(1), m.group(2).strip()
            lied = Lied(id=f"tpp-{nummer}-{slug(titel)}", titel=titel, referenties=[Referentie("tpp", nummer)], artiest="The Psalm Project")
        else:
            titel = str(tekst).strip()
            lied = Lied(id=f"tpp-{slug(titel)}", titel=titel, artiest="The Psalm Project")
        toegevoegd["tpp.yaml"] += _voeg_toe(catalogus, lied, "tpp.yaml")

    for (titel,) in _rijen(wb, "Bron SvG", (1,)):
        lied = Lied(id=f"svg-{slug(titel)}", titel=str(titel).strip(), artiest="Schrijvers voor Gerechtigheid")
        toegevoegd["svg.yaml"] += _voeg_toe(catalogus, lied, "svg.yaml")

    catalogus.schrijf()
    return toegevoegd


def _rijen(wb, blad: str, kolommen: tuple[int, ...]):
    if blad not in wb.sheetnames:
        return
    for i, rij in enumerate(wb[blad].iter_rows(values_only=True)):
        if i == 0:
            continue
        waarden = tuple(rij[k] if k < len(rij) else None for k in kolommen)
        if all(w is None or str(w).strip() == "" for w in waarden):
            continue
        yield waarden


def _voeg_toe(catalogus: Catalogus, lied: Lied, bestand: str) -> int:
    if lied.id in catalogus.liederen:
        return 0
    catalogus.voeg_lied_toe(lied, bestand)
    return 1
